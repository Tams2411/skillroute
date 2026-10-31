"""Resume parser with PDF/DOCX/TXT support."""

import logging
import sys
from functools import lru_cache
from pathlib import Path
from typing import Union

import pypdf
import pdfplumber
from docx import Document

from career_coach.config import settings

logger = logging.getLogger(__name__)


class ResumeParser:
    """Parse resume files (PDF, DOCX, TXT) into plain text."""

    SUPPORTED_FORMATS = {".pdf", ".docx", ".txt"}
    MAX_SIZE_BYTES = settings.max_pdf_size_mb * 1024 * 1024
    MAX_PAGES = settings.max_pdf_pages

    def __init__(self):
        self._cache: dict[str, str] = {}

    def extract_text(self, file_path: Union[str, Path]) -> str:
        """
        Extract text from a resume file with in-memory caching.

        Args:
            file_path: Path to the resume file

        Returns:
            Extracted text content

        Raises:
            ValueError: If file format is unsupported or too large
            FileNotFoundError: If file does not exist
        """
        path = Path(file_path)
        suffix = path.suffix.lower()

        if suffix not in self.SUPPORTED_FORMATS:
            raise ValueError(
                f"Unsupported file format: {suffix}. Supported: {self.SUPPORTED_FORMATS}"
            )

        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        # Check file size
        file_size = path.stat().st_size
        if file_size > self.MAX_SIZE_BYTES:
            raise ValueError(
                f"File too large: {file_size / (1024*1024):.1f}MB > {settings.max_pdf_size_mb}MB limit"
            )

        # Check cache
        cache_key = f"{path.name}:{file_size}:{path.stat().st_mtime}"
        if cache_key in self._cache:
            logger.info(f"Cache hit for: {path.name}")
            return self._cache[cache_key]

        logger.info(f"Extracting text from: {path.name}")

        if suffix == ".pdf":
            text = self._extract_pdf(path)
        elif suffix == ".docx":
            text = self._extract_docx(path)
        else:
            text = self._extract_txt(path)

        # Cache result
        self._cache[cache_key] = text
        return text

    def _extract_pdf(self, file_path: Path) -> str:
        """Extract text from PDF using pypdf with pdfplumber fallback."""
        errors = []

        # Try pypdf first
        try:
            extracted_parts = []
            with open(file_path, "rb") as f:
                reader = pypdf.PdfReader(f)

                # Check page count
                if len(reader.pages) > self.MAX_PAGES:
                    raise ValueError(
                        f"PDF has {len(reader.pages)} pages, limit is {self.MAX_PAGES}"
                    )

                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        extracted_parts.append(page_text)

            text = "\n".join(extracted_parts).strip()
            if text:
                return text
            errors.append("pypdf: no text extracted")
        except Exception as e:
            logger.debug(f"pypdf extraction failed: {e}")
            errors.append(f"pypdf: {str(e)}")

        # Fallback to pdfplumber
        try:
            with pdfplumber.open(file_path) as pdf:
                if len(pdf.pages) > self.MAX_PAGES:
                    raise ValueError(
                        f"PDF has {len(pdf.pages)} pages, limit is {self.MAX_PAGES}"
                    )

                text = "\n".join(page.extract_text() or "" for page in pdf.pages).strip()
                if text:
                    return text
                if not errors:
                    return ""
        except Exception as e:
            logger.debug(f"pdfplumber extraction failed: {e}")
            errors.append(f"pdfplumber: {str(e)}")

        if errors:
            raise Exception(f"PDF extraction failed: {'; '.join(errors)}")
        return ""

    def _extract_docx(self, file_path: Path) -> str:
        """Extract text from DOCX file including tables."""
        try:
            doc = Document(str(file_path))

            paragraphs = [para.text for para in doc.paragraphs]

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text)

            return "\n".join(p for p in paragraphs if p.strip()).strip()
        except Exception as e:
            logger.error(f"DOCX extraction failed: {e}")
            raise

    def _extract_txt(self, file_path: Path) -> str:
        """Extract text from plain text file."""
        try:
            return file_path.read_text(encoding="utf-8", errors="ignore").strip()
        except Exception as e:
            logger.error(f"TXT extraction failed: {e}")
            raise

    def clear_cache(self) -> None:
        """Clear the extraction cache."""
        self._cache.clear()


# Convenience function for simple usage
@lru_cache(maxsize=32)
def extract_resume_text_cached(file_path: str) -> str:
    """Cached wrapper for simple extraction."""
    parser = ResumeParser()
    return parser.extract_text(file_path)